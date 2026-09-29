"""Validate native table source edits with independent XML and PPTX readers."""
import argparse
import hashlib
import json
from pathlib import Path
from xml.sax.saxutils import escape
import zipfile
import jsonschema
from lxml import etree
from pptx import Presentation


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--corpus-dir', type=Path, required=True)
    parser.add_argument('--parity-dir', type=Path, required=True)
    parser.add_argument('--schema-dir', type=Path, required=True)
    parser.add_argument('--contracts-dir', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    assert not args.output.exists(), 'do not overwrite verification evidence'
    pml = etree.XMLSchema(etree.parse(str(args.schema_dir / 'pml.xsd')))
    dml = etree.XMLSchema(etree.parse(str(args.schema_dir / 'dml-main.xsd')))
    contracts = {kind: jsonschema.Draft202012Validator(json.loads((args.contracts_dir / f'{kind}.schema.json').read_text()))
                 for kind in ['pptx-source-response', 'pptx-import-request', 'pptx-import-response', 'pptx-text-edits']}
    manifest = json.loads((args.parity_dir / 'verification.json').read_bytes())
    rows = []
    for case in manifest['cases']:
        name = case['name']
        path = args.parity_dir / (name + '.pptx')
        source_path = args.corpus_dir / (name + '.pptx')
        assert hashlib.sha256(path.read_bytes()).hexdigest() == case['outputSha256']
        assert hashlib.sha256(source_path.read_bytes()).hexdigest() == case['sourceSha256']
        values = {suffix: json.loads((args.parity_dir / f'{name}.{suffix}.json').read_bytes())
                  for suffix in ['before', 'after', 'import-request', 'imported', 'edit-request']}
        for suffix, kind in [('before', 'pptx-source-response'), ('after', 'pptx-source-response'),
                             ('import-request', 'pptx-import-request'), ('imported', 'pptx-import-response'), ('edit-request', 'pptx-text-edits')]:
            contracts[kind].validate(values[suffix])
        request = values['edit-request']
        xsd_checks = 0
        with zipfile.ZipFile(source_path) as source, zipfile.ZipFile(path) as after:
            assert set(source.namelist()) == set(after.namelist())
            for part in source.namelist():
                original, actual = source.read(part), after.read(part)
                edits = [e for e in request['edits'] if e['target']['part'].lstrip('/') == part]
                if edits:
                    expected = original.decode('utf-8')
                    for edit in edits:
                        old = escape(edit['expectedText'])
                        assert expected.count(old) == 1, (name, part, 'fixture leaves must be unique')
                        expected = expected.replace(old, escape(edit['replacement']), 1)
                    assert expected.encode('utf-8') == actual, (name, part)
                else:
                    assert actual == original, (name, part)
                if part.endswith('.xml'):
                    node = etree.fromstring(actual)
                    if node.tag.startswith('{http://schemas.openxmlformats.org/presentationml/2006/main}'):
                        pml.assertValid(node)
                        xsd_checks += 1
                    if node.tag == '{http://schemas.openxmlformats.org/drawingml/2006/main}tblStyleLst':
                        dml.assertValid(node)
                        xsd_checks += 1
                    for table in node.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}tbl'):
                        dml.assertValid(table)
                        xsd_checks += 1
        tables = [Presentation(p).slides[0].shapes[0].table for p in [source_path, path]]
        before, after = tables
        assert [c.width for c in before.columns] == [c.width for c in after.columns]
        assert [r.height for r in before.rows] == [r.height for r in after.rows]
        edits = {(e['target']['paragraph'], e['target']['run']): e['replacement'] for e in request['edits']}
        flat = 0
        cells = 0
        for r in range(len(before.rows)):
            for c in range(len(before.columns)):
                old, new = before.cell(r, c), after.cell(r, c)
                assert (old.is_spanned, old.span_width, old.span_height) == (new.is_spanned, new.span_width, new.span_height)
                assert len(old.text_frame.paragraphs) == len(new.text_frame.paragraphs)
                for old_p, new_p in zip(old.text_frame.paragraphs, new.text_frame.paragraphs):
                    assert len(old_p.runs) == len(new_p.runs)
                    for i, (old_r, new_r) in enumerate(zip(old_p.runs, new_p.runs)):
                        assert new_r.text == edits.get((flat, i), old_r.text), (name, r, c, flat, i)
                    flat += 1
                cells += 1
        rows.append({'name': name, 'physicalCells': cells, 'editedLeaves': len(edits), 'xsdChecks': xsd_checks, 'jsonSchemaChecks': 5})
    report = {'scope': 'all other OPC part bytes and XML tokens preserved; independent table text, dimensions and merge readback; no application acceptance',
              'presentations': len(rows), 'physicalCells': sum(r['physicalCells'] for r in rows), 'editedLeaves': sum(r['editedLeaves'] for r in rows),
              'xsdChecks': sum(r['xsdChecks'] for r in rows), 'jsonSchemaChecks': sum(r['jsonSchemaChecks'] for r in rows), 'cases': rows}
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({k: v for k, v in report.items() if k != 'cases'}, ensure_ascii=False))


if __name__ == '__main__':
    main()
