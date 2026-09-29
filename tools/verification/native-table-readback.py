"""Independently read table exports after the native/WASM transaction sequence."""
import argparse
import hashlib
import json
from pathlib import Path
import zipfile
from lxml import etree
from pptx import Presentation


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--parity-dir', type=Path, required=True)
    parser.add_argument('--schema-dir', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    assert not args.output.exists(), 'do not overwrite verification evidence'
    pml = etree.XMLSchema(etree.parse(str(args.schema_dir / 'pml.xsd')))
    dml = etree.XMLSchema(etree.parse(str(args.schema_dir / 'dml-main.xsd')))
    source = (args.parity_dir / 'cases.json').read_bytes()
    rows = []
    for case in json.loads(source):
        if case['response']['status'] not in ['initialized', 'prepared']:
            continue
        name = 'original' if case['name'] == 'initialize' else case['name']
        path = args.parity_dir / (name + '.pptx')
        document = case['response']['snapshot']['document']
        expected = document['objects']['shape:1']['content']['table']
        schema_checks = 0
        with zipfile.ZipFile(path) as package:
            for part in package.namelist():
                if not part.endswith('.xml'):
                    continue
                node = etree.fromstring(package.read(part))
                if node.tag.startswith('{http://schemas.openxmlformats.org/presentationml/2006/main}'):
                    pml.assertValid(node)
                    schema_checks += 1
                for table in node.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}tbl'):
                    dml.assertValid(table)
                    schema_checks += 1
        presentation = Presentation(path)
        assert len(presentation.slides) == 1 and len(presentation.slides[0].shapes) == 1
        shape = presentation.slides[0].shapes[0]
        assert shape.has_table
        actual = shape.table
        assert len(actual.columns) == len(expected['columns']) and len(actual.rows) == len(expected['rows'])
        assert [c.width for c in actual.columns] == [int(c['width']) for c in expected['columns']]
        assert [r.height for r in actual.rows] == [int(r['height']) for r in expected['rows']]
        assert shape.width == sum(c.width for c in actual.columns)
        assert shape.height == sum(r.height for r in actual.rows)
        cells = 0
        for r, row in enumerate(expected['rows']):
            for c, cell in enumerate(row['cells']):
                cells += 1
                current, merge, body = actual.cell(r, c), cell['merge'], cell.get('text')
                text = '' if body is None else '\n'.join(''.join(run['content']['text'] for run in p['runs']) for p in body['paragraphs'])
                assert current.text == text, (name, r, c)
                assert current.is_spanned == (merge['kind'] == 'covered'), (name, r, c)
                if merge['kind'] == 'span':
                    assert (current.span_height, current.span_width) == (merge['rows'], merge['columns']), (name, r, c)
        rows.append({'name': name, 'physicalCells': cells, 'xsdChecks': schema_checks, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
    assert len(rows) == 11
    report = {
        'scope': 'independent structural readback after transactions; not application editing or rendering',
        'presentations': len(rows), 'physicalCells': sum(r['physicalCells'] for r in rows),
        'xsdChecks': sum(r['xsdChecks'] for r in rows), 'cases': rows,
        'inputSha256': hashlib.sha256(source).hexdigest(),
        'scriptSha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
    }
    args.output.write_text(json.dumps(report, indent=2) + '\n')
    print(json.dumps({k: report[k] for k in ['presentations', 'physicalCells', 'xsdChecks']}))


if __name__ == '__main__':
    main()
