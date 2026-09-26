"""Independent actual-file validation via ECMA XSD, lxml and python-pptx.

Schemas must be supplied explicitly from the official ECMA archive; no network
lookups occur here. This does not claim Office/WPS acceptance or kernel rendering.
"""
import argparse
import hashlib
import json
from pathlib import Path
import posixpath
import zipfile

from lxml import etree
from pptx import Presentation

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"p": P, "a": A, "r": R}


def digest(data):
    return hashlib.sha256(data).hexdigest()


def rich_text(body):
    paragraphs = []
    for p in body.findall(f"{{{A}}}p"):
        runs = []
        for run in p:
            if run.tag == f"{{{A}}}br":
                runs.append("\v")
            elif run.tag in (f"{{{A}}}r", f"{{{A}}}fld"):
                runs.append(run.findtext(f"{{{A}}}t", ""))
        paragraphs.append("".join(runs))
    return paragraphs


def author_text(body):
    return ["".join(run["content"].get("text", "\v" if run["content"]["kind"] == "break" else "\t") for run in p["runs"]) for p in body["paragraphs"]]


def native_geometry(node, expected):
    geometry = expected["content"].get("geometry")
    if not geometry:
        return
    sp = node.find("p:spPr", NS)
    if geometry["kind"] != "path":
        native = sp.find("a:prstGeom", NS)
        assert native.get("prst") == {"rectangle": "rect", "ellipse": "ellipse", "roundRectangle": "roundRect"}[geometry["kind"]]
        if geometry["kind"] == "roundRectangle":
            adjustment = int(native.find("a:avLst/a:gd", NS).get("fmla").removeprefix("val "))
            short = min(int(v) for v in expected["transform"]["size"].values())
            assert adjustment * short == int(geometry["radius"]) * 100000
    else:
        path = sp.find("a:custGeom/a:pathLst/a:path", NS)
        assert path.get("w") == geometry["viewport"]["width"]
        assert path.get("h") == geometry["viewport"]["height"]
        assert len(path) == len(geometry["commands"])
        for actual, command in zip(path, geometry["commands"]):
            tag, fields = {"move": ("moveTo", ["to"]), "line": ("lnTo", ["to"]), "quadratic": ("quadBezTo", ["control", "to"]), "cubic": ("cubicBezTo", ["control1", "control2", "to"]), "close": ("close", [])}[command["kind"]]
            assert etree.QName(actual).localname == tag
            assert [dict(point.attrib) for point in actual] == [command[field] for field in fields]


def native_body_style(node, body):
    props = node.find("p:txBody/a:lstStyle/a:defPPr/a:defRPr", NS)
    style = body["style"]
    if style.get("size", {}).get("kind") == "value":
        assert int(props.get("sz")) * 127 == int(style["size"]["value"])
    if style.get("color", {}).get("kind") == "value":
        expected = style["color"]["value"]
        if expected["kind"] == "srgb":
            rgb = expected["rgba"]
            actual = props.find("a:solidFill/a:srgbClr", NS)
            assert actual.get("val") == "".join(f"{rgb[c]:02X}" for c in ["red", "green", "blue"])
    modes = [p["style"].get("direction", {}).get("value", "leftToRight") for p in body["paragraphs"]]
    vert = "horz"
    if modes and modes[0] in ["verticalLeftToRight", "verticalRightToLeft"]:
        vert = {"verticalLeftToRight": "mongolianVert", "verticalRightToLeft": "eaVert"}[modes[0]]
    assert node.find("p:txBody/a:bodyPr", NS).get("vert", "horz") == vert
    paragraphs = node.findall("p:txBody/a:p", NS)
    for actual, expected in zip(paragraphs, body["paragraphs"]):
        direction = expected["style"].get("direction", {}).get("value")
        if direction in ["leftToRight", "rightToLeft"]:
            assert (actual.find("a:pPr", NS).get("rtl", "0") in ["1", "true"]) == (direction == "rightToLeft")


def native_appearance(node, expected):
    tag = "p:grpSpPr" if expected["content"]["kind"] == "group" else "p:spPr"
    props = node.find(tag, NS)
    fill = expected["appearance"].get("fill", {})
    if fill.get("kind") != "value":
        return
    fill = fill["value"]
    if fill["kind"] == "none":
        assert props.find("a:noFill", NS) is not None
    elif fill["color"]["kind"] == "srgb":
        rgba = fill["color"]["rgba"]
        actual = props.find("a:solidFill/a:srgbClr", NS)
        assert actual is not None
        assert actual.get("val") == "".join(f"{rgba[c]:02X}" for c in ["red", "green", "blue"])
        alpha = actual.find("a:alpha", NS)
        opacity = 100000 if alpha is None else int(alpha.get("val"))
        # Compare decoded 8-bit opacity, independently from the writer's rounding.
        assert round(opacity * 255 / 100000) == rgba["alpha"]


def native_stroke(node, expected):
    declaration=expected['appearance'].get('stroke',{})
    if declaration.get('kind')!='value':return
    s=declaration['value'];tag='p:grpSpPr' if expected['content']['kind']=='group' else 'p:spPr'
    line=node.find(tag+'/a:ln',NS);assert line is not None
    if s['kind']=='none':
        assert line.find('a:noFill',NS) is not None
        return
    assert int(line.get('w'))==int(s['width'])
    assert line.get('cap')=={'flat':'flat','round':'rnd','square':'sq'}.get(s.get('cap'))
    joins=[e for e in line if etree.QName(e).localname in ['round','bevel','miter']]
    if s.get('join') is None:assert not joins
    else:
        assert len(joins)==1 and etree.QName(joins[0]).localname==s['join']['kind']
        if s['join']['kind']=='miter':
            actual=joins[0].get('lim');expected=s['join'].get('limit')
            assert (None if actual is None else int(actual))==expected
    if s['color']['kind']=='srgb':
        rgba=s['color']['rgba'];c=line.find('a:solidFill/a:srgbClr',NS);assert c is not None
        assert c.get('val')==''.join(f'{rgba[k]:02X}' for k in ['red','green','blue'])
        alpha=c.find('a:alpha',NS);value=100000 if alpha is None else int(alpha.get('val'))
        assert round(value*255/100000)==rgba['alpha']


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--request", type=Path, required=True)
    parser.add_argument("--xsd-directory", type=Path, required=True)
    parser.add_argument("--external-roundtrip", action="store_true")
    args = parser.parse_args()
    request = json.loads(args.request.read_text())
    document = request["document"]
    safe_xml = etree.XMLParser(resolve_entities=False, no_network=True)
    validators = {P: etree.XMLSchema(etree.parse(str(args.xsd_directory / "pml.xsd"), safe_xml)), A: etree.XMLSchema(etree.parse(str(args.xsd_directory / "dml-main.xsd"), safe_xml))}
    native_objects = {}
    checked = []
    differences = []
    with zipfile.ZipFile(args.pptx) as package:
        assert package.testzip() is None, "invalid ZIP CRC"
        for name in package.namelist():
            if not name.endswith(".xml"):
                continue
            root = etree.fromstring(package.read(name), safe_xml)
            namespace = etree.QName(root).namespace
            if namespace in validators:
                # Other producers can have MCE extensions; full preprocessing is a
                # separate reader feature, so this gate applies to our generated output.
                if not args.external_roundtrip:
                    validators[namespace].assertValid(root)
                    checked.append(name)
                for node in root.xpath("//p:sp|//p:pic|//p:grpSp|//p:cxnSp", namespaces=NS):
                    nv = node.xpath("./*/p:cNvPr", namespaces=NS)
                    if not nv:
                        continue
                    object_name = nv[0].get("name")
                    assert object_name not in native_objects, f"duplicate named object {object_name}"
                    native_objects[object_name] = (name, node, nv[0])

        for id, expected in document["objects"].items():
            assert id in native_objects, f"missing object {id}"
            name, node, nv = native_objects[id]
            kind = expected["content"]["kind"]
            assert etree.QName(node).localname == {"shape": "sp", "picture": "pic", "group": "grpSp", "connector": "cxnSp"}[kind], id
            transform = node.find("p:grpSpPr/a:xfrm" if kind == "group" else "p:spPr/a:xfrm", NS)
            assert transform is not None, id
            for child, native, author in [("off", ["x", "y"], expected["transform"]["origin"]), ("ext", ["cx", "cy"], dict(zip(["cx", "cy"], expected["transform"]["size"].values())))]:
                for attr in native:
                    observed = int(transform.find(f"a:{child}", NS).get(attr))
                    if args.external_roundtrip and observed != int(author[attr]):
                        differences.append({"object": id, "property": f"transform.{child}.{attr}", "expected": author[attr], "actual": str(observed), "deltaEmu": str(observed-int(author[attr]))})
                    else:
                        assert observed == int(author[attr]), (id, child, attr, observed, author[attr])
            actual_rotation = int(transform.get("rot", "0")) % 21600000
            expected_rotation = expected["transform"]["rotation"] % 21600000
            if args.external_roundtrip and actual_rotation != expected_rotation:
                differences.append({"object": id, "property": "transform.rotation", "expected": str(expected_rotation), "actual": str(actual_rotation)})
            else:
                assert actual_rotation == expected_rotation, id
            body = expected["content"].get("text")
            if not args.external_roundtrip:
                native_geometry(node, expected)
                native_appearance(node, expected)
                native_stroke(node, expected)
                for attr, field in [("flipH", "flipHorizontal"), ("flipV", "flipVertical")]:
                    assert (transform.get(attr, "0") in ["1", "true"]) == expected["transform"][field]
            if body:
                actual = node.find("p:txBody", NS)
                assert actual is not None
                assert rich_text(actual) == author_text(body), (id, rich_text(actual), author_text(body))
                if not args.external_roundtrip:
                    native_body_style(node, body)
            if kind == "picture":
                relation = node.find("p:blipFill/a:blip", NS).get(f"{{{R}}}embed")
                folder, file = posixpath.split(name)
                relationships = etree.fromstring(package.read(f"{folder}/_rels/{file}.rels"), safe_xml)
                target = next(e.get("Target") for e in relationships if e.get("Id") == relation)
                resolved = posixpath.normpath(posixpath.join(folder, target)).lstrip("/")
                assert digest(package.read(resolved)) == document["resources"][expected["content"]["resource"]]["sha256"]
            if kind == "connector":
                for endpoint, tag in [("start", "stCxn"), ("end", "endCxn")]:
                    binding = expected["content"][endpoint]
                    if binding["kind"] == "attached":
                        native = node.find(f"p:nvCxnSpPr/p:cNvCxnSpPr/a:{tag}", NS)
                        assert native is not None
                        assert native.get("id") == native_objects[binding["object"]][2].get("id")
                        assert int(native.get("idx")) == binding["site"]
            if kind == "group":
                names = [c.xpath("./*/p:cNvPr", namespaces=NS)[0].get("name") for c in node if etree.QName(c).localname in ["sp", "pic", "grpSp", "cxnSp"]]
                assert names == expected["content"]["children"], id

    parsed = Presentation(args.pptx)
    assert len(parsed.slides) == len(document["slideOrder"])
    assert parsed.slide_width == int(document["pageSize"]["width"])
    assert parsed.slide_height == int(document["pageSize"]["height"])
    for i, id in enumerate(document["slideOrder"]):
        assert [s.name for s in parsed.slides[i].shapes] == document["slides"][id]["objects"]
        assert (parsed.slides[i]._element.get("show", "1") in ["0", "false"]) == document["slides"][id]["hidden"]
    checks = ["ZIP CRC", "native object kinds", "order and group ownership", "geometry and rotation", "Unicode paragraph text", "image content digest", "connector references"]
    if not args.external_roundtrip:
        checks += ["geometry commands", "explicit RGB fill and decoded 8-bit opacity", "native editable stroke width/cap/join/color", "explicit text body size/color", "writing mode and explicit paragraph direction", "flips", "hidden slides"]
    print(json.dumps({"format": "musteroffice.pptx-independent/1", "result": "differences" if differences else "passed", "fileSha256": digest(args.pptx.read_bytes()), "requestSha256": digest(args.request.read_bytes()), "externalRoundtrip": args.external_roundtrip, "schemasChecked": checked, "slides": len(parsed.slides), "nativeObjectsCompared": len(document["objects"]), "checks": checks, "geometryToleranceEmu": 0, "differences": differences, "scope": "Generated subset only. External geometry differences are reported without an invented acceptance tolerance. No Office/WPS, playback or kernel rendering claim."}, indent=2))


if __name__ == "__main__":
    main()
