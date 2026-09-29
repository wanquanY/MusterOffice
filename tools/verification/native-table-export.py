"""Check the owned 3x3 native table export; no Office/rendering acceptance claim."""
import argparse
import hashlib
import importlib.metadata
import json
from pathlib import Path
import subprocess
import zipfile
from lxml import etree
from pptx import Presentation


def digest(data):
    return hashlib.sha256(data).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cli', type=Path, required=True)
    parser.add_argument('--request', type=Path, required=True)
    parser.add_argument('--schema-dir', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=False)
    empty = args.output / 'empty.bin'
    empty.write_bytes(b'')
    output = args.output / 'author.pptx'
    command = [str(args.cli.resolve()), 'pptx-export', str(args.request.resolve()), str(empty), str(output)]
    run = subprocess.run(command, capture_output=True)
    (args.output / 'stdout.json').write_bytes(run.stdout)
    (args.output / 'stderr.log').write_bytes(run.stderr)
    run.check_returncode()
    pml = etree.XMLSchema(etree.parse(str(args.schema_dir / 'pml.xsd')))
    dml = etree.XMLSchema(etree.parse(str(args.schema_dir / 'dml-main.xsd')))
    checked = []
    with zipfile.ZipFile(output) as package:
        for name in package.namelist():
            if not name.endswith('.xml'):
                continue
            node = etree.fromstring(package.read(name))
            if node.tag.startswith('{http://schemas.openxmlformats.org/presentationml/2006/main}'):
                pml.assertValid(node)
                checked.append(name)
            for table in node.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}tbl'):
                dml.assertValid(table)
                checked.append(name + '#table')
    presentation = Presentation(output)
    assert len(presentation.slides) == 1
    slide = presentation.slides[0]
    assert len(slide.shapes) == 1 and slide.shapes[0].has_table
    table = slide.shapes[0].table
    assert len(table.rows) == len(table.columns) == 3
    assert table.cell(0, 0).is_merge_origin
    assert (table.cell(0, 0).span_height, table.cell(0, 0).span_width) == (2, 2)
    for r, c in [(0, 1), (1, 0), (1, 1)]:
        assert table.cell(r, c).is_spanned
    for r in range(3):
        for c in range(3):
            expected = '' if (r, c) == (2, 2) else f'cell {r}:{c} 中文 & <🚀> é'
            assert table.cell(r, c).text == expected
    assert table.cell(0, 0).margin_left == 10000
    report = {
        'scope': 'native table export structure only',
        'inputSha256': digest(args.request.read_bytes()),
        'cliSha256': digest(args.cli.read_bytes()),
        'scriptSha256': digest(Path(__file__).read_bytes()),
        'dependencies': {n: importlib.metadata.version(n) for n in ['lxml', 'python-pptx']},
        'schemaSha256': {p.name: digest(p.read_bytes()) for p in sorted(args.schema_dir.glob('*.xsd'))},
        'xsdParts': checked,
        'pythonPptxReadback': {'slides': 1, 'tables': 1, 'physicalCells': 9, 'mergedRectangle': [0, 0, 2, 2], 'textCells': 8},
        'sha256': digest(output.read_bytes()),
        'bytes': output.stat().st_size,
        'notVerified': ['table reader', 'table rendering', 'WPS', 'PowerPoint', 'WASM', 'product integration', 'performance', 'installer size'],
    }
    (args.output / 'verification.json').write_text(json.dumps(report, indent=2) + '\n')
    print(json.dumps({k: report[k] for k in ['scope', 'xsdParts', 'sha256', 'bytes']}))


if __name__ == '__main__':
    main()
