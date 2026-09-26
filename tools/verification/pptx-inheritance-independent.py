"""Verify relationships and effective bounds using ZIP/lxml/python-pptx.

MCE-selected XML is supplied to python-pptx only in an ephemeral in-memory copy.
That library does not process AlternateContent itself. No projected copy is used
as an edit/export result. Source preservation is checked by the separate oracle.
"""
import argparse
from functools import lru_cache
import hashlib
from io import BytesIO
import json
from pathlib import Path, PurePosixPath
import posixpath
import zipfile
from lxml import etree
from pptx import Presentation
from mce_reference import project, P, A, R

REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"p": P, "a": A}


def verify(case):
    index = case["response"]["index"]
    with zipfile.ZipFile(case["source"]) as package:
        raw = {n: package.read(n) for n in package.namelist()}
    logical = dict(raw)
    for part in [index["mainPart"], *index["surfaces"]]:
        logical[part.lstrip("/")] = etree.tostring(project(raw[part.lstrip("/")])[0])

    def part_ref(part):
        return {"part": part, "sha256": hashlib.sha256(raw[part.lstrip("/")]).hexdigest()}

    links = {}
    for part, surface in index["surfaces"].items():
        path = PurePosixPath(part)
        relpath = str(path.parent / "_rels" / (path.name + ".rels")).lstrip("/")
        resolved = {"layout": None, "master": None, "theme": None, "themeOverride": None}
        if relpath in raw:
            for rel in etree.fromstring(raw[relpath]):
                if rel.get("TargetMode") == "External":
                    continue
                kind = rel.get("Type").removeprefix(R + "/")
                target = posixpath.normpath(posixpath.join(str(path.parent), rel.get("Target")))
                if kind == "slideLayout" and surface["kind"] == "slide":
                    resolved["layout"] = target
                elif kind == "slideMaster":
                    resolved["master"] = target
                elif kind in ["theme", "themeOverride"]:
                    resolved[kind] = part_ref(target)
        assert resolved == surface["links"], (case["name"], part)
        links[part] = resolved

    @lru_cache(None)
    def theme(part):
        link = links[part]; parent = link["layout"] or link["master"]
        result = theme(parent) if parent else {"base": None, "overrides": []}
        return {"base": link["theme"] or result["base"], "overrides": result["overrides"] + ([link["themeOverride"]] if link["themeOverride"] else [])}

    for part, surface in index["surfaces"].items():
        assert surface["effectiveTheme"] == theme(part)

    stream = BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as package:
        for name, value in logical.items():
            package.writestr(name, value)
    stream.seek(0)
    presentation = Presentation(stream)
    surfaces = {str(slide.part.partname): slide for slide in presentation.slides}
    for master in presentation.slide_masters:
        surfaces[str(master.part.partname)] = master
        for layout in master.slide_layouts:
            surfaces[str(layout.part.partname)] = layout

    compared = 0; skipped = []
    for part, surface in index["surfaces"].items():
        # These names identify deliberate ambiguity/context probes. python-pptx
        # chooses the first duplicate, so cannot serve as their independent oracle.
        ambiguous = case["name"] in ["inspect-ambiguous-layout", "inspect-ambiguous-master"]
        native_objects = {o["nativeId"]: o for o in surface["objects"]}

        def walk(shapes):
            nonlocal compared
            for shape in shapes:
                native = native_objects[shape.shape_id]
                if hasattr(shape, "shapes"):
                    walk(shape.shapes)
                ph = shape.element.find(".//p:nvPr/p:ph", NS)
                if ph is not None:
                    expected = {"kind": ph.get("type"), "index": int(ph.get("idx")) if ph.get("idx") is not None else None,
                                "orientation": ph.get("orient"), "size": ph.get("sz"),
                                "customPrompt": None if ph.get("hasCustomPrompt") is None else ph.get("hasCustomPrompt") in ["true", "1"]}
                    assert expected == native["placeholder"]
                else:
                    assert native["placeholder"] is None
                if (ambiguous and ph is not None) or native["resolution"]["placeholderMatch"]["status"] == "unsupportedContext":
                    skipped.append({"part": part, "object": shape.shape_id, "reason": "ambiguity/context checked by explicit corpus assertions; third-party library lacks matching behavior"})
                    continue
                actual = native["resolution"]
                for field, attrs in [("origin", ["left", "top"]), ("size", ["width", "height"])]:
                    expected = tuple(getattr(shape, a) for a in attrs)
                    resolved = actual[field]
                    observed = (None, None) if resolved is None else tuple(int(resolved["value"][a]) for a in (["x", "y"] if field == "origin" else ["width", "height"]))
                    assert observed == expected, (case["name"], part, shape.name, field, observed, expected)
                    if resolved is not None:
                        declaration = shape
                        # Find the explicit author rather than inheriting a
                        # flattened effective value back into a model record.
                        attr = "x" if field == "origin" else "cx"
                        while getattr(declaration.element, attr) is None:
                            declaration = declaration._base_placeholder
                            assert declaration is not None
                        assert resolved["declaredBy"] == {"part": str(declaration.part.partname), "nativeId": declaration.shape_id}
                    compared += 1
        walk(surfaces[part].shapes)
    return {"name": case["name"], "sourceSha256": index["sourceSha256"], "surfaces": len(links),
            "geometryPropertiesCompared": compared, "libraryExclusions": skipped}


def main():
    parser = argparse.ArgumentParser(); parser.add_argument("parity_report", type=Path)
    args = parser.parse_args(); report = json.loads(args.parity_report.read_text())
    # Focus on dedicated hierarchy fixtures; the complete source corpus is also
    # checked for source preservation and object/text projection separately.
    selected = {"inheritance-chain", "inheritance-visible", "explicit-zero", "type-differs-index-matches", "unmatched", "detached", "unsupported-context",
                "ambiguous-layout", "ambiguous-master", "title-family", "mce-placeholder-parent", "theme-override-chain"}
    cases = [verify(c) for c in report["cases"] if c["name"].removeprefix("inspect-") in selected and "response" in c]
    print(json.dumps({"format": "musteroffice.inheritance-independent/1", "passed": len(cases), "cases": cases,
                      "scope": "Actual source relationships/theme stack and placeholder bounds; no styles, world coordinates or Office/WPS acceptance."}, indent=2))


if __name__ == "__main__":
    main()
